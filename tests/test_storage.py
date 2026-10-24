import datetime as dt

import pytest
from openpyxl import Workbook, load_workbook

from money_flow.lookups import LookupDataError, load_options
from money_flow.models import Options, Transaction, ValidationError, parse_amount, parse_date
from money_flow.storage import HEADERS, ExcelStorage, SchemaError


def make_txn(amount: float = 12.5, **kw) -> Transaction:
    defaults = dict(
        date=dt.date(2026, 10, 12), amount=amount, currency="GBP",
        payment_method="Revolut", category="Groceries", description="",
    )
    defaults.update(kw)
    return Transaction(**defaults)


# --- parsing ---------------------------------------------------------------
@pytest.mark.parametrize("text,expected", [
    ("12.5", 12.5), ("-2500", -2500.0), ("12,50", 12.5), ("1,234.56", 1234.56), (" 3 ", 3.0),
])
def test_parse_amount_ok(text, expected):
    assert parse_amount(text) == expected


@pytest.mark.parametrize("text", ["", "abc", "0", "nan", "inf", "1.2.3"])
def test_parse_amount_rejects(text):
    with pytest.raises(ValidationError):
        parse_amount(text)


def test_parse_date():
    assert parse_date("2026-01-31") == dt.date(2026, 1, 31)
    with pytest.raises(ValidationError):
        parse_date("31/01/2026")


def test_unknown_category_rejected():
    opts = Options(currencies=("GBP",), payment_methods=("Cash",), categories=("Food",))
    with pytest.raises(ValidationError) as exc:
        Transaction.from_form(
            options=opts, date="2026-10-12", amount="5",
            currency="GBP", payment_method="Cash", category="Nope",
        )
    assert exc.value.field == "category"


def test_from_form_ok():
    opts = Options(currencies=("GBP",), payment_methods=("Cash",), categories=("Food",))
    txn = Transaction.from_form(
        options=opts, date="2026-10-12", amount="-5,5",
        currency="GBP", payment_method="Cash", category="Food", description="  hi ",
    )
    assert (txn.amount, txn.kind, txn.description) == (-5.5, "Income", "hi")


# --- lookups ---------------------------------------------------------------
def _workbook_with_lookups(path, currencies, payments, categories):
    wb = Workbook()
    wb.active.title = "Transactions"
    for name, values in (("Currencies", currencies), ("Payment_methods", payments),
                         ("Categories", categories)):
        if values is None:
            continue
        ws = wb.create_sheet(name)
        for v in values:
            ws.append([v])
    wb.save(path)


def test_load_options_reads_column_a(tmp_path):
    path = tmp_path / "budget.xlsx"
    _workbook_with_lookups(
        path,
        ["GBP", "EUR"],
        ["Cash", None, "  Card  "],              # blank skipped, whitespace trimmed
        ["Food", "food", "Rent", 2024],          # case-insensitive dedupe, numbers -> str
    )
    opts = load_options(path)
    assert opts.currencies == ("GBP", "EUR")
    assert opts.payment_methods == ("Cash", "Card")
    assert opts.categories == ("Food", "Rent", "2024")


def test_load_options_reports_all_problems(tmp_path):
    path = tmp_path / "budget.xlsx"
    _workbook_with_lookups(path, ["GBP"], None, [])
    with pytest.raises(LookupDataError) as exc:
        load_options(path)
    msg = str(exc.value)
    assert "'Payment_methods' is missing" in msg
    assert "'Categories' has no values" in msg


def test_new_workbook_gets_empty_lookup_sheets(tmp_path):
    store = ExcelStorage(tmp_path / "budget.xlsx")
    store.ensure_exists()
    assert load_workbook(store.path).sheetnames == [
        "Transactions", "Currencies", "Payment_methods", "Categories",
    ]
    with pytest.raises(LookupDataError):
        load_options(store.path)


def test_append_preserves_lookup_sheets(tmp_path):
    path = tmp_path / "budget.xlsx"
    store = ExcelStorage(path)
    store.ensure_exists()
    wb = load_workbook(path)
    wb["Categories"].append(["Groceries"])
    wb.save(path)

    store.append(make_txn())
    assert load_workbook(path)["Categories"]["A1"].value == "Groceries"


# --- storage ---------------------------------------------------------------
def test_creates_file_with_headers(tmp_path):
    store = ExcelStorage(tmp_path / "sub" / "budget.xlsx")
    assert store.ensure_exists() is True
    ws = load_workbook(store.path)["Transactions"]
    assert tuple(c.value for c in ws[1]) == HEADERS
    assert ws.freeze_panes == "A2"
    assert store.ensure_exists() is False


def test_append_preserves_history(tmp_path):
    store = ExcelStorage(tmp_path / "budget.xlsx")
    assert store.append(make_txn(10)) == 2
    # Same date -> inserted at the top (row 2), pushing the first entry down to row 3
    assert store.append(make_txn(-2000, category="Salary", description="Oct pay")) == 2

    df = store.read_all()
    assert list(df.columns) == list(HEADERS)
    assert df["Amount"].tolist() == [-2000.0, 10.0]
    assert df["Date"].iloc[0] == dt.date(2026, 10, 12)
    assert df["Description"].tolist() == ["Oct pay", ""]


def test_append_keeps_dates_descending(tmp_path):
    store = ExcelStorage(tmp_path / "budget.xlsx")
    # Insert out of order: middle date, oldest date, newest date, then another middle date
    assert store.append(make_txn(10, date=dt.date(2026, 10, 15), description="mid-1")) == 2
    assert store.append(make_txn(20, date=dt.date(2026, 10, 10), description="oldest")) == 3
    assert store.append(make_txn(30, date=dt.date(2026, 10, 20), description="newest")) == 2
    assert store.append(make_txn(40, date=dt.date(2026, 10, 15), description="mid-2")) == 3

    df = store.read_all()
    assert df["Date"].tolist() == [
        dt.date(2026, 10, 20),
        dt.date(2026, 10, 15),
        dt.date(2026, 10, 15),
        dt.date(2026, 10, 10),
    ]
    assert df["Description"].tolist() == ["newest", "mid-2", "mid-1", "oldest"]

    ws = load_workbook(store.path)["Transactions"]
    assert ws.cell(row=2, column=1).number_format == "yyyy-mm-dd"
    assert ws.cell(row=5, column=2).number_format == "#,##0.00"


def test_schema_mismatch_detected(tmp_path):
    path = tmp_path / "budget.xlsx"
    wb = Workbook()
    wb.active.title = "Transactions"
    wb.active.append(["Wrong", "Headers"])
    wb.save(path)
    with pytest.raises(SchemaError):
        ExcelStorage(path).append(make_txn())

