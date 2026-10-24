"""Load dropdown options (currencies, payment methods, categories) from the workbook.

Each lookup sheet lists one option per row in column A, starting at row 1.
Sheet names come from ``config``; the values themselves live only in your
private ``budget_data.xlsx``, never in the source code.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from . import config
from .models import Options
from .storage import FileLockedError, StorageError


class LookupDataError(StorageError):
    """A lookup sheet is missing or empty."""


def load_options(path: Path | str = config.DATA_FILE) -> Options:
    """Read all lookup sheets and return them as an :class:`Options`.

    Raises :class:`LookupDataError` listing *every* problem at once (missing or
    empty sheets) so the user can fix the workbook in one go.
    """
    path = Path(path)
    if not path.exists():
        raise LookupDataError(f"Workbook not found: {path}")

    try:
        # read_only: fast, and works even while Excel has the file open.
        wb = load_workbook(path, read_only=True, data_only=True)
    except PermissionError as exc:
        raise FileLockedError(
            f"Could not read '{path.name}' - it is locked by another program."
        ) from exc
    except Exception as exc:  # corrupt / not a real xlsx
        raise StorageError(f"Could not open '{path}':\n{exc}") from exc

    try:
        sheetnames = list(wb.sheetnames)
        wanted = {
            "currencies": config.CURRENCIES_SHEET,
            "payment_methods": config.PAYMENT_METHODS_SHEET,
            "categories": config.CATEGORIES_SHEET,
        }
        values: dict[str, tuple[str, ...]] = {}
        problems: list[str] = []
        for key, sheet in wanted.items():
            if sheet not in sheetnames:
                problems.append(f"• Sheet '{sheet}' is missing.")
                continue
            items = _read_column_a(wb[sheet])
            if not items:
                problems.append(f"• Sheet '{sheet}' has no values in column A.")
            values[key] = items
    finally:
        wb.close()  # read-only workbooks keep the file handle open until closed

    if problems:
        raise LookupDataError(
            f"Cannot load dropdown options from '{path.name}':\n"
            + "\n".join(problems)
            + f"\n\nSheets found: {', '.join(sheetnames)}"
            + "\nEach lookup sheet must list one option per row in column A."
        )
    return Options(**values)


def _read_column_a(ws) -> tuple[str, ...]:
    """Non-blank values of column A, trimmed, de-duplicated (case-insensitive), in order."""
    seen: set[str] = set()
    result: list[str] = []
    for (cell,) in ws.iter_rows(min_col=1, max_col=1, values_only=True):
        if cell is None:
            continue
        if isinstance(cell, float) and cell.is_integer():
            cell = int(cell)  # 2024.0 -> "2024"
        text = str(cell).strip()
        if text and text.casefold() not in seen:
            seen.add(text.casefold())
            result.append(text)
    return tuple(result)
