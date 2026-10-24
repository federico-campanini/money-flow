"""Excel-backed storage for transactions.

Design notes
------------
* **Appends use openpyxl directly** (not ``DataFrame.to_excel``) so the
  existing workbook - its formatting, column widths, and any extra sheets you
  add by hand - is preserved. Only one new row is written per save.
* **Reads use pandas** (``read_all``) which is the convenient format for the
  analysis/reporting features planned for later phases.
* **Saves are atomic**: the workbook is written to a temp file in the same
  folder and then swapped in with ``os.replace``. A crash mid-save can never
  leave a half-written ``budget_data.xlsx``.
* **Lock detection**: on Windows, Excel holds an exclusive lock on open files,
  so we probe for write access *before* doing any work and translate the
  ``PermissionError`` into a friendly :class:`FileLockedError`.
"""

from __future__ import annotations

import datetime as dt
import os
import sys
import tempfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import config
from .models import Transaction

# Column header -> width (in Excel character units). Order defines the schema.
COLUMNS: dict[str, int] = {
    "Date": 12,
    "Amount": 12,
    "Currency": 10,
    "Payment Method": 18,
    "Category": 18,
    "Description": 50,
}
HEADERS: tuple[str, ...] = tuple(COLUMNS)
_DATE_COL = HEADERS.index("Date") + 1
_AMOUNT_COL = HEADERS.index("Amount") + 1

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="305496")


class StorageError(Exception):
    """Base class for storage problems the GUI should report to the user."""


class FileLockedError(StorageError):
    """The workbook is open in another program (typically Excel)."""


class SchemaError(StorageError):
    """The workbook exists but its header row doesn't match the expected schema."""


class ExcelStorage:
    def __init__(self, path: Path | str = config.DATA_FILE, sheet_name: str = config.SHEET_NAME):
        self.path = Path(path)
        self.sheet_name = sheet_name

    # ------------------------------------------------------------------ #
    # Public API
    # ------------------------------------------------------------------ #
    def ensure_exists(self) -> bool:
        """Create the workbook with headers if missing. Returns True if created."""
        if self.path.exists():
            self._assert_writable()
            wb = self._load()
            if self.sheet_name not in wb.sheetnames:
                self._init_sheet(wb.create_sheet(self.sheet_name))
                self._save(wb)
            else:
                self._validate_header(wb[self.sheet_name])
            return False

        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = self.sheet_name
        self._init_sheet(ws)
        # Empty lookup sheets so a brand-new workbook shows where the dropdown
        # options go. (Existing workbooks are never modified this way.)
        for name in (config.CURRENCIES_SHEET, config.PAYMENT_METHODS_SHEET, config.CATEGORIES_SHEET):
            wb.create_sheet(name)
        self._save(wb)
        return True

    def append(self, txn: Transaction) -> int:
        """Insert a transaction in descending date order (newest at the top)
        and return the Excel row number it was written to.
        """
        self.ensure_exists()
        wb = self._load()
        ws = wb[self.sheet_name]
        self._validate_header(ws)

        last_row = self._last_data_row(ws)
        target_row = self._find_insertion_row(ws, txn.date, last_row)

        if target_row <= last_row:
            ws.insert_rows(target_row)
            last_row += 1
        else:
            last_row = target_row

        values = (
            txn.date,
            float(txn.amount),
            txn.currency,
            txn.payment_method,
            txn.category,
            txn.description or None,
        )
        for col, value in enumerate(values, start=1):
            ws.cell(row=target_row, column=col, value=value)
        ws.cell(row=target_row, column=_DATE_COL).number_format = config.EXCEL_DATE_FORMAT
        ws.cell(row=target_row, column=_AMOUNT_COL).number_format = config.EXCEL_AMOUNT_FORMAT

        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last_row}"
        self._save(wb)
        return target_row

    def read_all(self) -> pd.DataFrame:
        """Return all transactions as a DataFrame (empty if the file doesn't exist)."""
        if not self.path.exists():
            return pd.DataFrame(columns=list(HEADERS))
        try:
            df = pd.read_excel(self.path, sheet_name=self.sheet_name, engine="openpyxl")
        except PermissionError as exc:
            raise FileLockedError(self._locked_message()) from exc
        if not df.empty:
            df["Date"] = pd.to_datetime(df["Date"]).dt.date
            df["Description"] = df["Description"].fillna("")
        return df

    # ------------------------------------------------------------------ #
    # Internals
    # ------------------------------------------------------------------ #
    def _init_sheet(self, ws: Worksheet) -> None:
        ws.append(list(HEADERS))
        for idx, (name, width) in enumerate(COLUMNS.items(), start=1):
            cell = ws.cell(row=1, column=idx)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"

    def _validate_header(self, ws: Worksheet) -> None:
        found = tuple(
            ws.cell(row=1, column=c).value for c in range(1, len(HEADERS) + 1)
        )
        if found != HEADERS:
            raise SchemaError(
                f"Unexpected columns in '{self.path.name}' (sheet '{self.sheet_name}').\n"
                f"Expected: {', '.join(HEADERS)}\n"
                f"Found:    {', '.join(str(v) for v in found)}"
            )

    @staticmethod
    def _last_data_row(ws: Worksheet) -> int:
        """Last row containing any value (ignores formatted-but-empty trailing rows)."""
        for row in range(ws.max_row, 1, -1):
            if any(ws.cell(row=row, column=c).value is not None for c in range(1, len(HEADERS) + 1)):
                return row
        return 1  # only the header

    @classmethod
    def _find_insertion_row(cls, ws: Worksheet, txn_date: dt.date, last_row: int) -> int:
        """Find the row index to keep dates descending (most recent at the top).

        Same-date entries are placed at the top of that date's group so the
        most recently entered transaction appears first.
        """
        for row in range(2, last_row + 1):
            row_date = cls._coerce_date(ws.cell(row=row, column=_DATE_COL).value)
            if row_date is not None and txn_date >= row_date:
                return row
        return last_row + 1

    @staticmethod
    def _coerce_date(value) -> dt.date | None:
        if isinstance(value, dt.datetime):
            return value.date()
        if isinstance(value, dt.date):
            return value
        if isinstance(value, str):
            try:
                return dt.date.fromisoformat(value.strip()[:10])
            except ValueError:
                return None
        return None

    def _load(self):
        try:
            return load_workbook(self.path)
        except PermissionError as exc:
            raise FileLockedError(self._locked_message()) from exc
        except Exception as exc:  # corrupt / not a real xlsx
            raise StorageError(f"Could not open '{self.path}':\n{exc}") from exc

    def _save(self, wb: Workbook) -> None:
        fd, tmp_name = tempfile.mkstemp(
            prefix=f".{self.path.stem}-", suffix=".xlsx.tmp", dir=self.path.parent
        )
        os.close(fd)
        tmp = Path(tmp_name)
        try:
            wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError as exc:
            raise FileLockedError(self._locked_message()) from exc
        except OSError as exc:
            raise StorageError(f"Could not save '{self.path}':\n{exc}") from exc
        finally:
            tmp.unlink(missing_ok=True)

    def _assert_writable(self) -> None:
        """Fail fast if another process holds the file."""
        try:
            with open(self.path, "r+b"):
                pass
        except PermissionError as exc:
            raise FileLockedError(self._locked_message()) from exc

        # macOS/Linux have no mandatory locks; fall back to office lock files.
        if sys.platform != "win32":
            for lock in (f"~${self.path.name}", f".~lock.{self.path.name}#"):
                if (self.path.parent / lock).exists():
                    raise FileLockedError(self._locked_message())

    def _locked_message(self) -> str:
        return (
            f"'{self.path.name}' appears to be open in another program "
            "(e.g. Microsoft Excel).\n\nPlease close it and try again."
        )
